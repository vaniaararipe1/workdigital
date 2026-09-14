"""
Painel de preview do slide - renderiza (de forma aproximada) o fundo e os
textos do slide da proposta, atualizando em tempo real conforme o
formulario e' preenchido. Nao e' o render final (isso sai do LibreOffice
na hora de gerar); e' so um retrato rapido para orientar o preenchimento.
"""
from __future__ import annotations

import io
import tkinter as tk
from pathlib import Path

from PIL import Image, ImageTk
from pptx import Presentation

PRETO = "#0B0B0F"
CINZA_CLARO = "#B8B8C4"
ROXO = "#6025E1"
COR_PADRAO = "#222222"
EMU_POR_PT = 12700


class SlidePreview(tk.Frame):
    LARGURA = 440
    ALTURA = 248

    def __init__(self, master, **kw):
        super().__init__(master, bg=PRETO, **kw)

        tk.Label(
            self,
            text="Preview do slide",
            bg=PRETO,
            fg="#FFFFFF",
            font=("Segoe UI", 10, "bold"),
        ).pack(anchor="w", pady=(0, 6))

        moldura = tk.Frame(self, bg="#3A3A46", padx=2, pady=2)
        moldura.pack()
        self.canvas = tk.Canvas(
            moldura,
            width=self.LARGURA,
            height=self.ALTURA,
            bg="white",
            highlightthickness=0,
        )
        self.canvas.pack()

        nav = tk.Frame(self, bg=PRETO)
        nav.pack(fill="x", pady=(8, 0))
        self.botao_anterior = tk.Button(
            nav,
            text="◀ Anterior",
            command=self.pagina_anterior,
            bg="#2A2A33",
            fg="white",
            relief="flat",
            activebackground="#3A3A46",
            activeforeground="white",
        )
        self.botao_anterior.pack(side="left")
        self.label_pagina = tk.Label(
            nav, text="", bg=PRETO, fg=CINZA_CLARO, font=("Segoe UI", 8)
        )
        self.label_pagina.pack(side="left", expand=True)
        self.botao_proxima = tk.Button(
            nav,
            text="Próxima ▶",
            command=self.pagina_proxima,
            bg="#2A2A33",
            fg="white",
            relief="flat",
            activebackground="#3A3A46",
            activeforeground="white",
        )
        self.botao_proxima.pack(side="right")

        self.prs: Presentation | None = None
        self.schema: dict | None = None
        self.paginas: list[int] = []
        self.pagina_atual = 0
        self.valores: dict = {}
        self._fundo_cache: dict[int, ImageTk.PhotoImage | None] = {}
        self._imagens_vivas: list[ImageTk.PhotoImage] = []
        self._escala = 1.0

    # ------------------------------------------------------------------ #
    def carregar(self, template_path: Path, schema: dict) -> None:
        self.prs = Presentation(str(template_path))
        self.schema = schema
        self._escala = self.LARGURA / self.prs.slide_width

        mapa_por_slide: dict[int, dict[int, dict]] = {}
        paginas: list[int] = []
        for campo in schema["campos"]:
            slide_num = campo["slide"]
            if slide_num not in paginas:
                paginas.append(slide_num)
            if campo["tipo_campo"] != "image":
                mapa_por_slide.setdefault(slide_num, {})[campo["shape_id"]] = campo
        paginas.sort()

        self._mapa_por_slide = mapa_por_slide
        self.paginas = paginas
        self.pagina_atual = 0
        self.valores = {}
        self._fundo_cache = {
            (n - 1): self._extrair_fundo(self.prs.slides[n - 1]) for n in paginas
        }
        self._atualizar_navegacao()
        self._desenhar()

    def definir_valores(self, valores: dict) -> None:
        self.valores = valores
        self._desenhar()

    def pagina_anterior(self) -> None:
        if self.pagina_atual > 0:
            self.pagina_atual -= 1
            self._atualizar_navegacao()
            self._desenhar()

    def pagina_proxima(self) -> None:
        if self.pagina_atual < len(self.paginas) - 1:
            self.pagina_atual += 1
            self._atualizar_navegacao()
            self._desenhar()

    # ------------------------------------------------------------------ #
    def _atualizar_navegacao(self) -> None:
        total = len(self.paginas)
        pos = self.pagina_atual + 1
        slide_num = self.paginas[self.pagina_atual] if self.paginas else 0
        self.label_pagina.config(text=f"Página {pos} de {total}  (slide {slide_num})")
        self.botao_anterior.config(state="normal" if self.pagina_atual > 0 else "disabled")
        self.botao_proxima.config(
            state="normal" if self.pagina_atual < total - 1 else "disabled"
        )

    def _extrair_fundo(self, slide) -> ImageTk.PhotoImage | None:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                try:
                    blob = shape.image.blob
                    img = Image.open(io.BytesIO(blob)).convert("RGB")
                    img = img.resize((self.LARGURA, self.ALTURA), Image.LANCZOS)
                    return ImageTk.PhotoImage(img)
                except Exception:
                    return None
        return None

    @staticmethod
    def _estilo(shape) -> tuple[float, bool, str]:
        tamanho_pt, negrito, cor = 11.0, False, COR_PADRAO
        p = shape.text_frame.paragraphs[0]
        run = p.runs[0] if p.runs else None
        if run is not None:
            if run.font.size is not None:
                tamanho_pt = run.font.size.pt
            negrito = bool(run.font.bold)
            try:
                if run.font.color and run.font.color.type is not None and int(run.font.color.type) == 1:
                    cor = "#" + str(run.font.color.rgb)
            except Exception:
                pass
        return tamanho_pt, negrito, cor

    def _valor_formatado(self, campo: dict, shape) -> str:
        chave = campo["chave"]
        valor = self.valores.get(chave)
        tipo_campo = campo["tipo_campo"]
        texto_original = shape.text_frame.text

        if tipo_campo == "single":
            return valor if valor else texto_original

        if tipo_campo == "suffix":
            runs = shape.text_frame.paragraphs[0].runs
            rotulo = runs[0].text if runs else ""
            if valor:
                return f"{rotulo}{valor}"
            return texto_original

        if tipo_campo in ("multiline", "multiline_com_titulo"):
            if not valor:
                return texto_original
            separador = "\n\n" if campo.get("blank_between", True) else "\n"
            corpo = separador.join(valor)
            if tipo_campo == "multiline_com_titulo":
                titulo = texto_original.split("\n", 1)[0]
                return f"{titulo}\n\n{corpo}"
            return corpo

        return texto_original

    def _desenhar(self) -> None:
        self.canvas.delete("all")
        self._imagens_vivas.clear()
        if not self.prs or not self.paginas:
            return

        slide_num = self.paginas[self.pagina_atual]
        idx = slide_num - 1
        slide = self.prs.slides[idx]
        mapa = self._mapa_por_slide.get(slide_num, {})

        fundo = self._fundo_cache.get(idx)
        if fundo is not None:
            self.canvas.create_image(0, 0, image=fundo, anchor="nw")
            self._imagens_vivas.append(fundo)
        else:
            self.canvas.create_rectangle(
                0, 0, self.LARGURA, self.ALTURA, fill="white", outline=""
            )

        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            campo = mapa.get(shape.shape_id)
            texto = self._valor_formatado(campo, shape) if campo else shape.text_frame.text
            if not texto.strip():
                continue
            tamanho_pt, negrito, cor = self._estilo(shape)
            x = shape.left * self._escala
            y = shape.top * self._escala
            largura = max(10, shape.width * self._escala)
            tamanho_px = max(6, round(tamanho_pt * EMU_POR_PT * self._escala))
            fonte = ("Segoe UI", tamanho_px, "bold" if negrito else "normal")
            self.canvas.create_text(
                x, y, text=texto, anchor="nw", width=largura,
                font=fonte, fill=cor, justify="left",
            )

        for campo in self.schema["campos"]:
            if campo["tipo_campo"] == "image" and campo["slide"] == slide_num:
                self._desenhar_logo(campo)

    def _desenhar_logo(self, campo: dict) -> None:
        left = campo["left"] * self._escala
        top = campo["top"] * self._escala
        max_w = campo["max_width"] * self._escala
        max_h = campo["max_height"] * self._escala
        caminho = self.valores.get(campo["chave"])

        if caminho and Path(caminho).is_file():
            try:
                img = Image.open(caminho).convert("RGBA")
                img.thumbnail((max(1, int(max_w)), max(1, int(max_h))), Image.LANCZOS)
                foto = ImageTk.PhotoImage(img)
                self._imagens_vivas.append(foto)
                cx = left + (max_w - img.width) / 2
                cy = top + (max_h - img.height) / 2
                self.canvas.create_image(cx, cy, image=foto, anchor="nw")
                return
            except Exception:
                pass

        self.canvas.create_rectangle(
            left, top, left + max_w, top + max_h, outline="#CCCCCC", dash=(4, 2)
        )
        self.canvas.create_text(
            left + max_w / 2,
            top + max_h / 2,
            text="Logo do cliente\n(opcional)",
            fill="#AAAAAA",
            font=("Segoe UI", 8),
            justify="center",
        )
