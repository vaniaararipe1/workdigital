#!/usr/bin/env python3
"""
Gerador automático de propostas comerciais Work Digital.

Le um arquivo de dados (JSON) com as informacoes do cliente, preenche o
template correspondente (simples ou completa) SEM alterar nada do design
original, e exporta o resultado em .pptx e .pdf prontos para envio.

Uso:
    python3 scripts/gerar_proposta.py --tipo simples --dados dados.json
    python3 scripts/gerar_proposta.py --tipo completa --dados dados.json --saida propostas/

O nome dos arquivos gerados segue o padrao:
    "<Cliente> Modelo de proposta comercial <simples|completa> (dd-mm-yyyy).pptx/.pdf"
"""
from __future__ import annotations

import argparse
import copy
import datetime
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

def _root_dir() -> Path:
    """Raiz de dados do projeto: pasta do repo em modo normal, ou a pasta
    temporaria de extracao quando empacotado com PyInstaller (--onefile)."""
    meipass = getattr(sys, "_MEIPASS", None)
    if meipass:
        return Path(meipass)
    return Path(__file__).resolve().parent.parent


ROOT = _root_dir()
FONTS_DIR = ROOT / "fonts"
MAX_BYTES = 5 * 1024 * 1024  # 5 MB

TIPOS = {
    "simples": ROOT / "scripts" / "schema_simples.json",
    "completa": ROOT / "scripts" / "schema_completa.json",
}

FROZEN = bool(getattr(sys, "_MEIPASS", None))
# App instalado (executavel): salva em Documentos, uma pasta estavel e
# visivel para o usuario. Rodando a partir do repositorio (dev/CLI): mantem
# a pasta propostas/ do proprio projeto, como ja documentado no README.
PASTA_SAIDA_PADRAO = (
    Path.home() / "Documents" / "Propostas Work Digital" if FROZEN else ROOT / "propostas"
)


# --------------------------------------------------------------------------- #
# Fontes
# --------------------------------------------------------------------------- #
def instalar_fontes() -> None:
    """Instala as fontes do template (Nunito / Space Grotesk Medium) no
    sistema, para que a conversao para PDF renderize com a tipografia
    correta em vez de uma fonte substituta. Best-effort e multiplataforma:
    nunca deve interromper a geracao da proposta se falhar."""
    try:
        if sys.platform.startswith("win"):
            _instalar_fontes_windows()
        elif sys.platform == "darwin":
            _instalar_fontes_simples(Path.home() / "Library" / "Fonts", cache=False)
        else:
            _instalar_fontes_simples(Path.home() / ".fonts", cache=True)
    except Exception:
        pass


def _instalar_fontes_simples(destino: Path, cache: bool) -> None:
    destino.mkdir(parents=True, exist_ok=True)
    mudou = False
    for f in FONTS_DIR.glob("*.ttf"):
        alvo = destino / f.name
        if not alvo.exists():
            shutil.copy(f, alvo)
            mudou = True
    if mudou and cache:
        subprocess.run(
            ["fc-cache", "-f", str(destino)],
            check=False,
            capture_output=True,
        )


def _instalar_fontes_windows() -> None:
    import winreg  # disponivel apenas no Windows

    destino = Path.home() / "AppData" / "Local" / "Microsoft" / "Windows" / "Fonts"
    destino.mkdir(parents=True, exist_ok=True)
    chave = winreg.OpenKey(
        winreg.HKEY_CURRENT_USER,
        r"Software\Microsoft\Windows NT\CurrentVersion\Fonts",
        0,
        winreg.KEY_SET_VALUE,
    )
    try:
        for f in FONTS_DIR.glob("*.ttf"):
            alvo = destino / f.name
            if not alvo.exists():
                shutil.copy(f, alvo)
            winreg.SetValueEx(
                chave, f"{f.stem} (TrueType)", 0, winreg.REG_SZ, str(alvo)
            )
    finally:
        chave.Close()

    # Registra a fonte no GDI da sessao atual e avisa os apps abertos, para
    # ficar disponivel *imediatamente* (sem precisar reiniciar o Windows) —
    # inclusive para este mesmo processo, que ainda vai criar a janela do Tk.
    try:
        import ctypes

        for f in FONTS_DIR.glob("*.ttf"):
            ctypes.windll.gdi32.AddFontResourceW(str(destino / f.name))
        HWND_BROADCAST = 0xFFFF
        WM_FONTCHANGE = 0x001D
        ctypes.windll.user32.SendMessageW(HWND_BROADCAST, WM_FONTCHANGE, 0, 0)
    except Exception:
        pass


# --------------------------------------------------------------------------- #
# Manipulacao de texto preservando formatacao original
# --------------------------------------------------------------------------- #
def _find_shape(prs: Presentation, slide_index_1based: int, shape_id: int):
    slide = prs.slides[slide_index_1based - 1]
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise KeyError(
        f"shape id {shape_id} nao encontrado no slide {slide_index_1based} "
        "(o template pode ter sido alterado)"
    )


def _clone_run_format(origem, destino) -> None:
    destino.font.size = origem.font.size
    destino.font.bold = origem.font.bold
    destino.font.italic = origem.font.italic
    destino.font.name = origem.font.name
    try:
        if origem.font.color and origem.font.color.type is not None:
            if origem.font.color.type == 1:  # MSO_COLOR_TYPE.RGB
                destino.font.color.rgb = origem.font.color.rgb
            else:
                destino.font.color.theme_color = origem.font.color.theme_color
    except Exception:
        pass


def set_single(shape, texto: str) -> None:
    """Substitui o texto mantendo a formatacao do primeiro run."""
    p = shape.text_frame.paragraphs[0]
    if not p.runs:
        return
    p.runs[0].text = texto
    for extra in p.runs[1:]:
        extra.text = ""


def set_suffix(shape, texto: str, manter_runs: int = 1) -> None:
    """Mantem os N primeiros runs (rotulo em negrito) e substitui o
    restante do paragrafo pelo texto informado."""
    p = shape.text_frame.paragraphs[0]
    runs = p.runs
    if len(runs) <= manter_runs:
        if not runs:
            return
        novo = p.add_run()
        _clone_run_format(runs[manter_runs - 1], novo)
        novo.text = texto
        return
    runs[manter_runs].text = texto
    for extra in runs[manter_runs + 1 :]:
        extra.text = ""


def _paragraph_templates(paragraphs, skip_first: bool = False):
    """Retorna (paragrafo_com_conteudo, paragrafo_em_branco) usados como
    modelo de formatacao para reconstruir a lista de linhas."""
    candidatos = paragraphs[1:] if skip_first else paragraphs
    conteudo = next((p for p in candidatos if p.text.strip()), candidatos[0])
    branco = next(
        (p for p in candidatos if not p.text.strip() and p is not conteudo),
        None,
    )
    return conteudo._p, (branco._p if branco is not None else None)


def _set_multiline(shape, linhas: list[str], blank_between: bool, skip_first: bool) -> None:
    tf = shape.text_frame
    txBody = tf._txBody
    paragraphs = tf.paragraphs

    conteudo_tpl, branco_tpl = _paragraph_templates(paragraphs, skip_first=skip_first)
    conteudo_tpl = copy.deepcopy(conteudo_tpl)
    branco_tpl = copy.deepcopy(branco_tpl) if branco_tpl is not None else None

    titulo_el = copy.deepcopy(paragraphs[0]._p) if skip_first else None

    for p_el in list(txBody.findall(qn("a:p"))):
        txBody.remove(p_el)

    if titulo_el is not None:
        txBody.append(titulo_el)

    def montar_paragrafo(texto: str):
        p_el = copy.deepcopy(conteudo_tpl)
        runs = p_el.findall(qn("a:r"))
        for extra in runs[1:]:
            p_el.remove(extra)
        if runs:
            t_el = runs[0].find(qn("a:t"))
            t_el.text = texto
        return p_el

    for i, linha in enumerate(linhas):
        txBody.append(montar_paragrafo(linha))
        se_ultimo = i == len(linhas) - 1
        if blank_between and branco_tpl is not None and not se_ultimo:
            txBody.append(copy.deepcopy(branco_tpl))


def set_multiline(shape, linhas: list[str], blank_between: bool = True) -> None:
    _set_multiline(shape, linhas, blank_between, skip_first=False)


def set_multiline_com_titulo(shape, linhas: list[str], blank_between: bool = True) -> None:
    _set_multiline(shape, linhas, blank_between, skip_first=True)


def adicionar_imagem(slide, caminho: str, campo: dict) -> None:
    """Insere a logo do cliente em uma caixa (left/top/max_width/max_height,
    em EMU) definida no schema, mantendo a proporcao original da imagem e
    sem aumentar logos pequenas (evita pixelizacao)."""
    caminho_imagem = Path(caminho)
    if not caminho_imagem.is_file():
        raise DadosInvalidos(f"Arquivo de imagem nao encontrado: {caminho}")

    max_largura = campo["max_width"]
    max_altura = campo["max_height"]

    figura = slide.shapes.add_picture(str(caminho_imagem), campo["left"], campo["top"])
    escala = min(max_largura / figura.width, max_altura / figura.height, 1.0)
    nova_largura = int(figura.width * escala)
    nova_altura = int(figura.height * escala)
    figura.width = nova_largura
    figura.height = nova_altura
    figura.left = campo["left"] + (max_largura - nova_largura) // 2
    figura.top = campo["top"] + (max_altura - nova_altura) // 2


def adicionar_botao_link(slide, campo: dict, url: str) -> None:
    """Insere um botao roxo de destaque, com hyperlink, no slide — usado
    para o link do wireframe/previa do projeto."""
    if not url.strip():
        return
    if not url.lower().startswith(("http://", "https://")):
        url = "https://" + url.strip()

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        campo["left"],
        campo["top"],
        campo["width"],
        campo["height"],
    )
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x60, 0x25, 0xE1)
    shape.line.fill.background()
    try:
        shape.adjustments[0] = 0.5
    except (IndexError, AttributeError):
        pass

    tf = shape.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    margem = campo.get("margem_botao", 45720)
    tf.margin_left = tf.margin_right = Emu(margem)
    tf.margin_top = tf.margin_bottom = Emu(0)

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = campo.get("texto_botao", "Veja aqui o wireframe do seu projeto")
    run.font.size = Pt(campo.get("tamanho_fonte_botao", 13))
    run.font.bold = True
    run.font.name = "Nunito"
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    run.hyperlink.address = url


def _como_lista_de_linhas(valor) -> list[str]:
    """Aceita tanto uma lista de linhas quanto uma string unica (nesse caso
    vira uma lista de 1 item) — nunca itera caractere por caractere."""
    if isinstance(valor, str):
        return [valor] if valor.strip() else []
    return list(valor)


APLICADORES = {
    "single": lambda shape, valor, campo: set_single(shape, str(valor)),
    "suffix": lambda shape, valor, campo: set_suffix(
        shape, str(valor), manter_runs=campo.get("manter_runs", 1)
    ),
    "multiline": lambda shape, valor, campo: set_multiline(
        shape, _como_lista_de_linhas(valor), blank_between=campo.get("blank_between", True)
    ),
    "multiline_com_titulo": lambda shape, valor, campo: set_multiline_com_titulo(
        shape, _como_lista_de_linhas(valor), blank_between=campo.get("blank_between", True)
    ),
}


# --------------------------------------------------------------------------- #
# Geracao
# --------------------------------------------------------------------------- #
class DadosInvalidos(Exception):
    """Dados faltando ou invalidos para gerar a proposta."""


class LibreOfficeNaoEncontrado(Exception):
    """LibreOffice (soffice) nao foi encontrado no computador."""


def carregar_schema(tipo: str) -> dict:
    if tipo not in TIPOS:
        raise DadosInvalidos(f"Tipo invalido '{tipo}'. Use: {', '.join(TIPOS)}")
    with open(TIPOS[tipo], encoding="utf-8") as f:
        return json.load(f)


def validar_dados(schema: dict, dados: dict) -> dict:
    """Preenche os campos faltantes com o valor padrao do schema e valida
    obrigatoriedade. Retorna o dicionario de dados completo."""
    completos = dict(dados)
    faltando = []
    for campo in schema["campos"]:
        chave = campo["chave"]
        if chave not in completos or completos[chave] in (None, "", []):
            if "padrao" in campo:
                completos[chave] = campo["padrao"]
            elif campo.get("obrigatorio"):
                faltando.append(campo.get("descricao", chave))
    if faltando:
        raise DadosInvalidos(
            "Faltam campos obrigatorios: " + ", ".join(faltando)
        )
    if "cliente" not in completos or not completos["cliente"]:
        raise DadosInvalidos("O campo 'cliente' e obrigatorio (usado tambem no nome do arquivo).")
    return completos


def preencher_template(schema: dict, dados: dict) -> Presentation:
    template_path = ROOT / schema["template"]
    prs = Presentation(template_path)
    for campo in schema["campos"]:
        chave = campo["chave"]
        if chave not in dados or not dados[chave]:
            continue
        if campo["tipo_campo"] == "image":
            slide = prs.slides[campo["slide"] - 1]
            adicionar_imagem(slide, dados[chave], campo)
            continue
        if campo["tipo_campo"] == "botao":
            slide = prs.slides[campo["slide"] - 1]
            adicionar_botao_link(slide, campo, dados[chave])
            continue
        shape = _find_shape(prs, campo["slide"], campo["shape_id"])
        aplicar = APLICADORES[campo["tipo_campo"]]
        aplicar(shape, dados[chave], campo)
    return prs


def nome_base(cliente: str, rotulo: str, data: str | None) -> str:
    if not data:
        data = datetime.date.today().strftime("%d-%m-%Y")
    else:
        data = data.replace("/", "-")
    cliente_limpo = " ".join(cliente.strip().split())
    return f"{cliente_limpo} Modelo de proposta comercial {rotulo} ({data})"


def localizar_soffice() -> str | None:
    """Procura o executavel do LibreOffice no PATH e em locais tipicos de
    instalacao no Windows/Mac/Linux."""
    encontrado = shutil.which("soffice") or shutil.which("soffice.exe")
    if encontrado:
        return encontrado
    candidatos: list[str] = []
    if sys.platform.startswith("win"):
        candidatos = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
    elif sys.platform == "darwin":
        candidatos = ["/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    else:
        candidatos = ["/usr/bin/soffice", "/usr/lib/libreoffice/program/soffice"]
    for candidato in candidatos:
        if Path(candidato).exists():
            return candidato
    return None


def exportar_pdf(pptx_path: Path, saida_dir: Path) -> Path:
    soffice = localizar_soffice()
    if not soffice:
        raise LibreOfficeNaoEncontrado(
            "LibreOffice nao foi encontrado neste computador. Instale gratuitamente "
            "em https://www.libreoffice.org/download para gerar o PDF automaticamente "
            "(o PPTX ja foi gerado normalmente)."
        )
    instalar_fontes()
    with tempfile.TemporaryDirectory(prefix="lo_profile_") as perfil:
        cmd = [
            soffice,
            "--headless",
            "--norestore",
            f"-env:UserInstallation=file://{perfil}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(saida_dir),
            str(pptx_path),
        ]
        resultado = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
    pdf_path = saida_dir / (pptx_path.stem + ".pdf")
    if not pdf_path.exists():
        raise RuntimeError(
            "Falha ao gerar o PDF via LibreOffice:\n"
            f"stdout: {resultado.stdout}\nstderr: {resultado.stderr}"
        )
    return pdf_path


class ResultadoGeracao:
    def __init__(self, pptx_path: Path, pdf_path: Path | None, aviso: str | None):
        self.pptx_path = pptx_path
        self.pdf_path = pdf_path
        self.aviso = aviso


def gerar(tipo: str, dados: dict, saida_dir: Path) -> ResultadoGeracao:
    schema = carregar_schema(tipo)
    dados_completos = validar_dados(schema, dados)
    prs = preencher_template(schema, dados_completos)

    saida_dir.mkdir(parents=True, exist_ok=True)
    base = nome_base(dados_completos["cliente"], schema["rotulo"], dados_completos.get("data"))
    pptx_path = saida_dir / f"{base}.pptx"
    prs.save(pptx_path)

    pdf_path: Path | None = None
    aviso: str | None = None
    try:
        pdf_path = exportar_pdf(pptx_path, saida_dir)
    except (LibreOfficeNaoEncontrado, RuntimeError) as e:
        aviso = str(e)

    for caminho in (pptx_path, pdf_path):
        if caminho is None:
            continue
        tamanho = caminho.stat().st_size
        if tamanho > MAX_BYTES:
            print(
                f"[AVISO] {caminho.name} tem {tamanho / 1024 / 1024:.2f} MB, "
                "acima do limite de 5 MB.",
                file=sys.stderr,
            )

    return ResultadoGeracao(pptx_path, pdf_path, aviso)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--tipo", required=True, choices=list(TIPOS))
    parser.add_argument("--dados", required=True, help="Caminho para o JSON com os dados do cliente")
    parser.add_argument(
        "--saida",
        default=str(PASTA_SAIDA_PADRAO),
        help=f"Diretorio de saida (padrao: {PASTA_SAIDA_PADRAO})",
    )
    args = parser.parse_args()

    with open(args.dados, encoding="utf-8") as f:
        dados = json.load(f)

    try:
        resultado = gerar(args.tipo, dados, Path(args.saida))
    except DadosInvalidos as e:
        raise SystemExit(f"Erro: {e}")

    print("Proposta gerada com sucesso:")
    print(
        f"  PPTX: {resultado.pptx_path} "
        f"({resultado.pptx_path.stat().st_size / 1024:.0f} KB)"
    )
    if resultado.pdf_path:
        print(
            f"  PDF:  {resultado.pdf_path} "
            f"({resultado.pdf_path.stat().st_size / 1024:.0f} KB)"
        )
    if resultado.aviso:
        print(f"  [AVISO] {resultado.aviso}", file=sys.stderr)


if __name__ == "__main__":
    main()
